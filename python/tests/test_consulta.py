"""Retrato do console: lista do Bacen (uma linha por requisição do CCS) cruzada com os work items pelo ofício."""
import json
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from simba import consulta, lista_bacen


def _ccs(n: int) -> str:
    return f"CCS2026100500000{n}"


def _item(n: int, status: str, resposta: str = "", tipo_excecao: str = "") -> dict:
    dados = {"Destino": "002-PF", "Caso": f"{n:06d}", "DV": f"0{n}", "Tipo": "Corretora", "Banco": "XP", "Vara": "1",
             "Investigados": [{}, {}]}
    return {
        "sys_id": f"id-{n}",
        "name": f"SimbaValidador_JUDTASK{n}",
        "status": status,
        "attempts_count": "1",
        "response_content": resposta,
        "exception_type": tipo_excecao,
        "sys_updated_on": "2026-10-07 20:00:00",
        "request_content": json.dumps(dados),
    }


def _requisicao(n: int, limite: str) -> dict:
    return {"ccs": _ccs(n), "atendimento": f"002-PF-{n:06d}-0{n}", "jud": f"JUD{n}", "marca": "XP Corretora",
            "estado_jud": "Open", "limite": limite, "sistema": "SIMBA"}


ONTEM = (date.today() - timedelta(days=1)).isoformat()
AMANHA = (date.today() + timedelta(days=1)).isoformat()


def _retrato() -> dict:
    return {
        "baixado_em": "2026-10-07T21:00:00",
        "itens": {
            "JUDTASK1": _item(1, "success", "validado e gerado; transmitido"),
            "JUDTASK2": _item(2, "failure", "[validando] O Número da Vara não pode ter mais de 10", "business"),
            "JUDTASK3": _item(3, "failure", "[preparando_envio] Órgão 002 sem chave no CSV de chaves", "application"),
            "JUDTASK4": _item(4, "pending"),
            "JUDTASK5": _item(5, "pending"),
        },
        "tarefas": {
            "JUDTASK1": {"evidence_attachment": "abc", "parent.number": "JUD1", "parent.state": "3",
                         "parent.official_letter_number": _ccs(1)},
            "JUDTASK2": {"parent.official_letter_number": f" {_ccs(2).lower()} "},
            "JUDTASK3": {"parent.number": "JUD3", "parent.state": "2", "parent.official_letter_number": _ccs(3)},
            "JUDTASK4": {"parent.official_letter_number": _ccs(4)},
        },
    }


def _lista() -> list[dict]:
    # Requisições 1 a 4 têm work item; a 9 não. O work item 5 (JUD sem ofício) fica fora da lista.
    sta = {**_requisicao(8, ONTEM), "sistema": "STA"}
    return [_requisicao(1, ONTEM), _requisicao(2, AMANHA), _requisicao(3, ONTEM), _requisicao(4, AMANHA),
            _requisicao(9, ONTEM), sta]


def _linhas() -> dict:
    return {l.judtask or l.ccs: l for l in consulta.montar_linhas(_retrato(), _lista())}


def test_situacao_de_cada_requisicao() -> None:
    linhas = _linhas()

    sucesso = linhas["JUDTASK1"]
    assert (sucesso.validado, sucesso.transmitido, sucesso.cabine) == (True, True, True)
    assert sucesso.ccs == _ccs(1) and sucesso.jud == "JUD1" and sucesso.prazo == "Vencido"

    dados = linhas["JUDTASK2"]
    assert dados.situacao == "Falha" and dados.categoria_erro == "Dados/arquivos"
    assert dados.etapa_erro == "validando" and not dados.validado
    assert dados.erro == "O Número da Vara não pode ter mais de 10"
    assert dados.ccs == _ccs(2)  # ofício casado sem diferença de maiúsculas/espaços

    sem_chave = linhas["JUDTASK3"]
    assert sem_chave.categoria_erro == "Ação manual" and sem_chave.validado and not sem_chave.transmitido

    assert linhas["JUDTASK4"].situacao == "Pendente" and linhas["JUDTASK4"].prazo == "No prazo"


def test_requisicao_sem_work_item_e_work_item_fora_da_lista() -> None:
    linhas = _linhas()

    falta = linhas[_ccs(9)]
    assert falta.situacao == "Sem work item" and falta.na_lista and not falta.work_item
    assert falta.simba_code == "002-PF-000009-09" and falta.jud == "JUD9"
    assert not linhas["JUDTASK5"].na_lista


def test_resumo_da_lista() -> None:
    resumo = consulta.resumo(consulta.montar_linhas(_retrato(), _lista()))

    assert resumo["Requisições"] == 5 and resumo["Atendimentos"] == 5
    assert resumo["Com work item"] == 4 and resumo["Sem work item"] == 1
    assert resumo["Validados"] == 2 and resumo["Transmitidos"] == 1 and resumo["Cabine"] == 1
    assert resumo["Com erro"] == 2
    assert resumo["Vencidas sem transmissão"] == 2  # requisições 3 e 9 (a 1 já foi transmitida)
    assert resumo["Work items fora da lista"] == 1
    assert resumo["Requisições STA (fora do RPA)"] == 1  # a requisição 8 não entra nos números do RPA


def test_sem_lista_resume_os_work_items() -> None:
    resumo = consulta.resumo(consulta.montar_linhas(_retrato(), []))

    assert resumo["Requisições"] == 5 and resumo["Work items fora da lista"] == 0


def test_le_a_lista_do_bacen_pela_aba_com_num_ctrl_ccs(tmp_path: Path) -> None:
    planilha = tmp_path / "3290.xlsx"
    livro = Workbook()
    outra = livro.active
    outra.title = "Conciliação SN - JD"
    outra.append(["Official letter number", "Number"])
    outra.append([_ccs(7), "JUD7"])
    aba = livro.create_sheet("3290")
    aba.append(["NUM_CTRL_CCS", "NUM_CTRL_ENVIO", "JUD ServiceNow", "Marca JUD", "Estado JUD", "CPFCNPJ_PESQUISADO",
                "DTHR_LIMITE_RESPOSTA", "CD_SISTEMA_ENVIO"])
    aba.append([_ccs(1).lower(), "002-PF-000001-01", "JUD1", "XP Corretora", "Closed", "12345678900", "20261029235959",
                "Simba"])
    aba.append([_ccs(2), "002-PF-000002-02", "#N/A", "XP Corretora", "#N/A", "98765432100", "20260812235959", "STA"])
    aba.append([None, None, None, None, None, None, None, None])
    livro.save(planilha)

    requisicoes = lista_bacen.ler(planilha)

    assert requisicoes == [
        {"ccs": _ccs(1), "atendimento": "002-PF-000001-01", "jud": "JUD1", "marca": "XP CORRETORA",
         "estado_jud": "CLOSED", "limite": "2026-10-29", "sistema": "SIMBA"},
        {"ccs": _ccs(2), "atendimento": "002-PF-000002-02", "jud": "", "marca": "XP CORRETORA",
         "estado_jud": "#N/A", "limite": "2026-08-12", "sistema": "STA"},
    ]
    # O CPF/CNPJ pesquisado não é lido.
    assert all("12345678900" not in json.dumps(r) for r in requisicoes)


def test_exporta_abas_com_erros_claros(tmp_path: Path) -> None:
    linhas = consulta.montar_linhas(_retrato(), _lista())

    destino = consulta.exportar(linhas, tmp_path / "export.xlsx", "2026-10-07T21:00:00")

    livro = load_workbook(destino)
    assert livro.sheetnames == ["Resumo", "Lista", "Erros", "Sem work item", "Fora da lista"]
    erros = list(livro["Erros"].values)
    assert len(erros) == 3  # cabeçalho + 2 erros
    assert "O Número da Vara não pode ter mais de 10" in erros[1]
    assert len(list(livro["Lista"].values)) == 7  # cabeçalho + 6 requisições (Simba e STA)
