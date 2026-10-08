"""Retrato dos casos para o console: a lista do Bacen (fonte da verdade) cruzada com uma leitura do ServiceNow.

Cada requisição do CCS da lista é uma linha, casada com o work item da fila pelo número de controle do CCS =
official_letter_number da JUD (comparado sem diferença de maiúsculas/espaços). Requisições sem work item mostram o
que falta gerar; work items que não casam com a lista aparecem como "fora da lista".

Sem base de dados: o ServiceNow é a fonte de verdade e o retrato vai para um arquivo JSON que só muda quando o
usuário baixa de novo. A leitura é paginada (1.000 registros por requisição) e passa pelo limite de requisições.
Só entram work items criados a partir de ITENS_DESDE.

Situação de cada caso:
- work item da fila Simba Validador: pendente / em andamento / sucesso / falha, tentativas, erro;
- validado: work item com sucesso, já transmitido, ou falha numa etapa depois da validação;
- transmitido: comprovante no campo evidence_attachment da JUDTASK;
- Cabine: JUD (parent da JUDTASK) Concluída — o RPA da Cabine conclui a JUD só depois de confirmar.
"""
import json
import re
from dataclasses import asdict, dataclass, fields
from datetime import datetime
from pathlib import Path
from typing import Optional

from simba import config, lista_bacen
from simba.fila import TABELA as TABELA_ITENS
from simba.fila import filtro_desde
from simba.servicenow import TABELA_TAREFAS, ServiceNow

FILA = "Simba Validador"
CAMPOS_ITEM = "sys_id,name,status,attempts_count,response_content,exception_type,sys_updated_on,request_content"
CAMPOS_TAREFA = "number,evidence_attachment,parent.number,parent.state,parent.simba_code,parent.official_letter_number"
JUD_CONCLUIDA = "3"
# JUDTASKs por requisição na consulta das tarefas (numberIN...).
TAREFAS_POR_CONSULTA = 300
# Etapas depois da validação: falhar nelas significa que o caso já foi validado.
ETAPAS_POS_VALIDACAO = ("preparando_envio", "enviando", "registrando", "anexando")
_ETAPA = re.compile(r"^\[(?P<etapa>[a-z_]+)\] ")

SITUACOES = {
    "pending": "Pendente",
    "in_progress": "Em andamento",
    "success": "Sucesso",
    "failure": "Falha",
    "": "Sem status",
}
CATEGORIAS = {"business": "Dados/arquivos", "application": "Ação manual"}


@dataclass
class Linha:
    ccs: str
    na_lista: bool
    sistema: str
    limite: str
    prazo: str
    judtask: str
    simba_code: str
    work_item: str
    situacao: str
    validado: bool
    transmitido: bool
    cabine: bool
    etapa_erro: str
    categoria_erro: str
    erro: str
    tentativas: int
    jud: str
    oficio: str
    tipo: str
    instituicao: str
    processo: str
    vara: str
    investigados: int
    atualizado_em: str


def judtask_do_item(nome: str) -> str:
    """SimbaValidador_JUDTASK0000001 -> JUDTASK0000001."""
    return nome.rsplit("_", 1)[-1]


def baixar(sn: ServiceNow) -> dict:
    """Lê do SN os work items da fila (desde ITENS_DESDE) e as JUDTASKs deles, com a JUD."""
    itens = sn.listar_todos(TABELA_ITENS, f"work_queue.name={FILA}{filtro_desde()}^ORDERBYsys_created_on", CAMPOS_ITEM)
    # Um work item por JUDTASK: vale o mais recente (os itens vêm em ordem de criação).
    por_judtask = {judtask_do_item(i["name"]): i for i in itens}
    numeros = sorted(por_judtask)
    tarefas: dict[str, dict[str, str]] = {}
    for inicio in range(0, len(numeros), TAREFAS_POR_CONSULTA):
        lote = numeros[inicio : inicio + TAREFAS_POR_CONSULTA]
        for t in sn.listar(TABELA_TAREFAS, "numberIN" + ",".join(lote), CAMPOS_TAREFA, len(lote)):
            tarefas[t["number"]] = t
    return {
        "baixado_em": datetime.now().isoformat(timespec="seconds"),
        "itens_desde": config.ITENS_DESDE,
        "itens": por_judtask,
        "tarefas": tarefas,
    }


def montar_linhas(retrato: Optional[dict], requisicoes: list[dict[str, str]]) -> list[Linha]:
    """Uma linha por requisição da lista (com o work item que casa, se houver) e uma por work item fora da lista."""
    retrato = retrato or {"itens": {}, "tarefas": {}}
    casos = {}
    for judtask, item in retrato["itens"].items():
        tarefa = retrato["tarefas"].get(judtask, {})
        oficio = lista_bacen.normalizar(tarefa.get("parent.official_letter_number") or _json(item["request_content"]).get("Ofício"))
        casos.setdefault(oficio, []).append((judtask, item, tarefa))
    linhas, na_lista = [], set()
    for requisicao in requisicoes:
        encontrados = casos.get(requisicao["ccs"], [])
        if not encontrados:
            linhas.append(_sem_work_item(requisicao))
        for judtask, item, tarefa in encontrados:
            linhas.append(_linha(judtask, item, tarefa, requisicao))
            na_lista.add(judtask)
    for judtask, item in retrato["itens"].items():
        if judtask not in na_lista:
            linhas.append(_linha(judtask, item, retrato["tarefas"].get(judtask, {}), None))
    return linhas


def _sem_work_item(requisicao: dict[str, str]) -> Linha:
    return Linha(**{
        **{f.name: "" for f in fields(Linha)},
        "ccs": requisicao["ccs"],
        "na_lista": True,
        "sistema": requisicao.get("sistema", ""),
        "limite": requisicao["limite"],
        "prazo": lista_bacen.prazo(requisicao["limite"]),
        "simba_code": requisicao["atendimento"],
        "jud": requisicao["jud"],
        "situacao": "Sem work item",
        "validado": False,
        "transmitido": False,
        "cabine": False,
        "tentativas": 0,
        "investigados": 0,
    })


def _linha(judtask: str, item: dict, tarefa: dict, requisicao: Optional[dict[str, str]]) -> Linha:
    dados = _json(item["request_content"])
    status = item["status"]
    resposta = item["response_content"]
    etapa = (m["etapa"] if (m := _ETAPA.match(resposta)) else "") if status != "success" else ""
    transmitido = bool(tarefa.get("evidence_attachment"))
    if status == "failure":
        categoria = CATEGORIAS.get(item["exception_type"], "Técnico")
    elif status == "pending" and resposta:
        categoria = "Técnico (nova tentativa)"
    else:
        categoria = ""
    limite = requisicao["limite"] if requisicao else ""
    return Linha(
        ccs=requisicao["ccs"] if requisicao else "",
        na_lista=requisicao is not None,
        sistema=requisicao.get("sistema", "") if requisicao else lista_bacen.SIMBA,
        limite=limite,
        prazo=lista_bacen.prazo(limite),
        judtask=judtask,
        simba_code=tarefa.get("parent.simba_code") or _atendimento(dados),
        work_item=item["sys_id"],
        situacao=SITUACOES.get(status, status),
        validado=status == "success" or transmitido or etapa in ETAPAS_POS_VALIDACAO,
        transmitido=transmitido,
        cabine=tarefa.get("parent.state") == JUD_CONCLUIDA,
        etapa_erro=etapa,
        categoria_erro=categoria,
        erro=_ETAPA.sub("", resposta) if categoria else "",
        tentativas=int(item["attempts_count"] or 0),
        jud=tarefa.get("parent.number", ""),
        oficio=tarefa.get("parent.official_letter_number", ""),
        tipo=dados.get("Tipo", ""),
        instituicao=dados.get("Banco", ""),
        processo=dados.get("Processo", ""),
        vara=str(dados.get("Vara", "")),
        investigados=len(dados.get("Investigados") or []),
        atualizado_em=item["sys_updated_on"],
    )


def _json(texto: str) -> dict:
    try:
        dados = json.loads(texto)
    except ValueError:
        return {}
    return dados if isinstance(dados, dict) else {}


def _atendimento(dados: dict) -> str:
    partes = [str(dados.get(k, "")).strip() for k in ("Destino", "Caso", "DV")]
    return "-".join(partes) if all(partes) else ""


def resumo(linhas: list[Linha]) -> dict[str, int]:
    """Números da lista do Bacen enviados pelo Simba (ou de todos os work items, sem lista) e a conferência."""
    tem_lista = any(l.na_lista for l in linhas)
    base = [l for l in linhas if l.na_lista and l.sistema == lista_bacen.SIMBA] if tem_lista else linhas
    com_item = [l for l in base if l.work_item]
    return {
        "Requisições": len({l.ccs for l in base}) if tem_lista else len(base),
        "Atendimentos": len({l.simba_code for l in base if l.simba_code}),
        "Com work item": len(com_item),
        "Sem work item": sum(not l.work_item for l in base),
        "Pendentes": sum(l.situacao == "Pendente" for l in com_item),
        "Em andamento": sum(l.situacao == "Em andamento" for l in com_item),
        "Validados": sum(l.validado for l in com_item),
        "Transmitidos": sum(l.transmitido for l in com_item),
        "Cabine": sum(l.cabine for l in com_item),
        "Com erro": sum(bool(l.categoria_erro) for l in com_item),
        "Vencidas sem transmissão": sum(l.prazo == "Vencido" and not l.transmitido for l in base),
        "Work items fora da lista": sum(not l.na_lista for l in linhas) if tem_lista else 0,
        "Requisições STA (fora do RPA)": sum(l.na_lista and l.sistema != lista_bacen.SIMBA for l in linhas),
    }


def salvar(retrato: dict, arquivo: Path) -> None:
    arquivo.parent.mkdir(parents=True, exist_ok=True)
    temporario = arquivo.with_suffix(".tmp")
    temporario.write_text(json.dumps(retrato, ensure_ascii=False), encoding="utf-8")
    temporario.replace(arquivo)


def carregar(arquivo: Path) -> Optional[dict]:
    return json.loads(arquivo.read_text(encoding="utf-8")) if arquivo.is_file() else None


TITULOS = {
    "ccs": "Requisição CCS",
    "na_lista": "Na lista",
    "sistema": "Sistema de envio",
    "limite": "Prazo de resposta",
    "prazo": "Situação do prazo",
    "judtask": "JUDTASK",
    "simba_code": "Atendimento Simba",
    "work_item": "Work item (sys_id)",
    "situacao": "Situação",
    "validado": "Validado",
    "transmitido": "Transmitido",
    "cabine": "Cabine",
    "etapa_erro": "Etapa do erro",
    "categoria_erro": "Categoria do erro",
    "erro": "Erro",
    "tentativas": "Tentativas",
    "jud": "JUD",
    "oficio": "Ofício (JUD)",
    "tipo": "Tipo",
    "instituicao": "Instituição",
    "processo": "Processo",
    "vara": "Vara",
    "investigados": "Investigados",
    "atualizado_em": "Atualizado em (UTC)",
}


def exportar(linhas: list[Linha], destino: Path, baixado_em: str) -> Path:
    """Excel com Resumo, Lista (todas as requisições), Erros, Sem work item e Fora da lista."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    aba = livro.active
    aba.title = "Resumo"
    aba.append(["Retrato do ServiceNow baixado em", baixado_em])
    for nome, valor in resumo(linhas).items():
        aba.append([nome, valor])
    colunas = [f.name for f in fields(Linha)]
    abas = {
        "Lista": [l for l in linhas if l.na_lista],
        "Erros": [l for l in linhas if l.categoria_erro],
        "Sem work item": [l for l in linhas if l.na_lista and not l.work_item],
        "Fora da lista": [l for l in linhas if not l.na_lista],
    }
    for titulo, grupo in abas.items():
        aba = livro.create_sheet(titulo)
        aba.append([TITULOS[c] for c in colunas])
        for celula in aba[1]:
            celula.font = Font(bold=True)
        for linha in grupo:
            valores = asdict(linha)
            aba.append([("Sim" if v else "Não") if isinstance(v, bool) else v for v in (valores[c] for c in colunas)])
        aba.auto_filter.ref = aba.dimensions
        aba.freeze_panes = "A2"
        for i, c in enumerate(colunas, 1):
            aba.column_dimensions[get_column_letter(i)].width = 80 if c == "erro" else max(12, len(TITULOS[c]) + 2)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


def arquivo_padrao() -> Path:
    return config.CONSOLE_DIR / "retrato.json"
