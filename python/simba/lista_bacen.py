"""Lista do Bacen (fonte da verdade do que precisa ser respondido): uma linha por requisição do CCS.

Planilha fixa (SIMBA_LISTA_BACEN), aba com a coluna NUM_CTRL_CCS (a aba "3290"). Cada requisição casa com a JUD do
ServiceNow pelo número de controle do CCS = official_letter_number da JUD. A lista não muda: o resultado da leitura
fica num cache JSON ao lado do retrato, refeito só quando a planilha muda. O CPF/CNPJ pesquisado não é lido.
"""
import json
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from simba import config

COLUNA_CCS = "NUM_CTRL_CCS"
COLUNAS = {
    "ccs": COLUNA_CCS,
    "atendimento": "NUM_CTRL_ENVIO",
    "jud": "JUD SERVICENOW",
    "marca": "MARCA JUD",
    "estado_jud": "ESTADO JUD",
    "limite": "DTHR_LIMITE_RESPOSTA",
    # SIMBA ou STA: este RPA responde só as requisições enviadas pelo Simba.
    "sistema": "CD_SISTEMA_ENVIO",
}
SIMBA = "SIMBA"


def normalizar(valor: object) -> str:
    """Texto para comparação: sem diferença de maiúsculas/espaços; 417618.0 do Excel vira 417618."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split()).upper()


def _data_limite(valor: object) -> str:
    """20261029235959 -> 2026-10-29 (vazio se não der para ler)."""
    texto = normalizar(valor)
    try:
        return datetime.strptime(texto[:8], "%Y%m%d").date().isoformat()
    except ValueError:
        return ""


def ler(planilha: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    livro = load_workbook(planilha, read_only=True, data_only=True)
    try:
        for aba in livro.worksheets:
            linhas = aba.iter_rows(values_only=True)
            cabecalho = [normalizar(c) for c in next(linhas, ())]
            if COLUNA_CCS not in cabecalho:
                continue
            indices = {campo: cabecalho.index(nome) for campo, nome in COLUNAS.items() if nome in cabecalho}
            requisicoes = []
            for linha in linhas:
                valores = {campo: normalizar(linha[i]) if i < len(linha) else "" for campo, i in indices.items()}
                if not valores.get("ccs"):
                    continue
                valores["jud"] = "" if valores.get("jud") == "#N/A" else valores.get("jud", "")
                valores["limite"] = _data_limite(valores.get("limite"))
                requisicoes.append({campo: valores.get(campo, "") for campo in COLUNAS})
            return requisicoes
    finally:
        livro.close()
    raise ValueError(f"Nenhuma aba da planilha tem a coluna {COLUNA_CCS}: {planilha}")


def carregar(planilha: Optional[Path] = None) -> Optional[dict]:
    """{"arquivo", "lida_em", "requisicoes"} da lista (do cache quando a planilha não mudou); None sem lista."""
    planilha = planilha or config.LISTA_BACEN
    if not planilha.is_file():
        return None
    # Muda com a planilha e com as colunas lidas (cache de versão anterior do código é refeito).
    assinatura = f"{planilha.resolve()}|{planilha.stat().st_size}|{planilha.stat().st_mtime_ns}|{','.join(COLUNAS)}"
    cache = config.CONSOLE_DIR / "lista_bacen.json"
    if cache.is_file():
        dados = json.loads(cache.read_text(encoding="utf-8"))
        if dados.get("assinatura") == assinatura:
            return dados
    dados = {
        "assinatura": assinatura,
        "arquivo": str(planilha),
        "lida_em": datetime.now().isoformat(timespec="seconds"),
        "requisicoes": ler(planilha),
    }
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
    return dados


def prazo(limite: str, hoje: Optional[date] = None) -> str:
    if not limite:
        return ""
    return "Vencido" if date.fromisoformat(limite) < (hoje or date.today()) else "No prazo"
